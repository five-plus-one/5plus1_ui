from pathlib import Path
import json,html
from docx import Document
from pptx import Presentation
from pptx.util import Inches,Pt
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
p=Path('public/samples');p.mkdir(parents=True,exist_ok=True)
titles=['阅读与创造','共同的设计语言','从想法到体验']
paragraphs=['让内容成为主角，让每一次行动都有明确的方向。','用留白组织信息，用清晰的文字连接想法。','从阅读到理解，从一个想法到完整的体验。']
doc=Document();doc.add_heading('共同的设计语言',0)
for i,t in enumerate(titles):
 doc.add_heading(t,1);doc.add_paragraph(paragraphs[i]);doc.add_paragraph('清晰、自然、有温度。五加一设计体系。')
doc.save(p/'design.docx')
prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(6.75)
for i,t in enumerate(titles):
 slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=t;slide.placeholders[1].text=paragraphs[i]+'\n\n五加一 Design'
prs.save(p/'design.pptx')
w=Workbook();ws=w.active;ws.title='组件清单';ws.append(['组件','用途','状态'])
for row in [('按钮','主要行动','可使用'),('文件预览','阅读资料','可使用'),('文件树','组织资料','可使用'),('协作编辑','一起创造','可使用')]:ws.append(row)
w.save(p/'design.xlsx')
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'));c=canvas.Canvas(str(p/'design.pdf'),pagesize=(595,842))
for i,t in enumerate(titles):
 c.setFillColorRGB(.21,.39,.28);c.setFont('STSong-Light',28);c.drawString(65,720,t);c.setFont('STSong-Light',14);c.setFillColorRGB(.25,.3,.25);c.drawString(65,665,paragraphs[i]);c.drawString(65,625,'清晰、自然、有温度。五加一设计体系。');c.setFont('Helvetica',10);c.drawString(65,60,f'FIVE PLUS ONE / DESIGN   {i+1:02}');c.showPage()
c.save()
for kind in ['pdf','docx','pptx','xlsx']:
 pages=[]
 for i,t in enumerate(titles):
  wide=kind=='pptx';width,height=(1200,675) if wide else (794,1123);x=75;top=120
  lines=''.join(f'<rect x="{x}" y="{top+245+j*34}" width="{width-2*x-(j%3)*45}" height="7" rx="3" fill="#e3e7df"/>' for j in range(7))
  svg=f'''<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#fffefa"/><text x="{x}" y="{top-45}" font-family="sans-serif" font-size="13" letter-spacing="3" fill="#6b7869">FIVE PLUS ONE / DESIGN NOTES</text><line x1="{x}" y1="{top-15}" x2="{width-x}" y2="{top-15}" stroke="#dfe5d9"/><text x="{x}" y="{top+65}" font-family="Microsoft YaHei,sans-serif" font-size="38" fill="#294a33">{t}</text><text x="{x}" y="{top+120}" font-family="Microsoft YaHei,sans-serif" font-size="17" fill="#687166">{paragraphs[i]}</text><rect x="{x}" y="{top+160}" width="{width-2*x}" height="48" rx="8" fill="#e7eedf"/><text x="{x+20}" y="{top+191}" font-family="Microsoft YaHei,sans-serif" font-size="16" fill="#356348">清晰、自然、有温度。</text>{lines if not wide else ''}<text x="{x}" y="{height-45}" font-family="sans-serif" font-size="12" fill="#687166">{kind.upper()} / DESIGN SPECIMEN</text><text x="{width-x-20}" y="{height-45}" font-family="sans-serif" font-size="12" fill="#687166">0{i+1}</text></svg>'''
  file=f'{kind}-page-{i+1}.svg';(p/file).write_text(svg,encoding='utf-8');pages.append({'n':i+1,'url':'/samples/'+file})
 (p/f'{kind}.json').write_text(json.dumps({'status':'ready','contentHash':'design-v1','manifest':{'pageCount':3,'pages':pages}}),encoding='utf-8')
print('Created PDF, DOCX, PPTX, XLSX and preview specimens.')
